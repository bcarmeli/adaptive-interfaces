from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


BG_DARK = RGBColor(0x16, 0x22, 0x2A)
BG_LIGHT = RGBColor(0xF3, 0xEF, 0xE8)
TEXT_DARK = RGBColor(0x18, 0x1A, 0x1C)
TEXT_LIGHT = RGBColor(0xF8, 0xF5, 0xF0)
ACCENT = RGBColor(0xC8, 0x5A, 0x38)
MUTED = RGBColor(0x6D, 0x73, 0x77)
CARD = RGBColor(0xFF, 0xFC, 0xF7)


def add_textbox(slide, x, y, w, h, text, size, color, bold=False, font="Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return box


def add_bullet_list(slide, x, y, w, h, items, size=18, color=TEXT_DARK):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        run = p.runs[0]
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_round_box(slide, x, y, w, h, fill, line_fill=None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line_fill or fill
    return shape


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "Language as Interface to Latent Space"
prs.core_properties.author = "OpenAI Codex"

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG_LIGHT

banner = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.95))
banner.fill.solid()
banner.fill.fore_color.rgb = ACCENT
banner.line.color.rgb = ACCENT

add_textbox(slide, Inches(0.65), Inches(1.15), Inches(8.4), Inches(0.8), "Language as an Interface to a Latent Space", 27, TEXT_DARK, bold=True, font="Aptos Display")
add_textbox(slide, Inches(0.68), Inches(1.95), Inches(8.1), Inches(0.82), "Communication is not the latent space itself. It is the public, addressable surface through which one agent gains partial access to another agent's hidden capabilities.", 17, MUTED)

add_round_box(slide, Inches(0.75), Inches(3.0), Inches(3.25), Inches(2.4), CARD)
add_textbox(slide, Inches(1.0), Inches(3.25), Inches(2.5), Inches(0.4), "Latent Space", 22, TEXT_DARK, bold=True, font="Aptos Display")
add_textbox(slide, Inches(1.0), Inches(3.75), Inches(2.6), Inches(1.3), "Private reasoning
Hidden heuristics
Internal capabilities", 18, TEXT_DARK)

add_round_box(slide, Inches(5.05), Inches(3.0), Inches(3.25), Inches(2.4), RGBColor(0xF4, 0xD8, 0xCD))
add_textbox(slide, Inches(5.3), Inches(3.25), Inches(2.6), Inches(0.4), "Interface", 22, TEXT_DARK, bold=True, font="Aptos Display")
add_textbox(slide, Inches(5.3), Inches(3.75), Inches(2.6), Inches(1.3), "Words, symbols,
APIs, examples,
public handles", 18, TEXT_DARK)

add_round_box(slide, Inches(9.35), Inches(3.0), Inches(3.25), Inches(2.4), RGBColor(0xD8, 0xE6, 0xD8))
add_textbox(slide, Inches(9.6), Inches(3.25), Inches(2.6), Inches(0.4), "Affordance", 22, TEXT_DARK, bold=True, font="Aptos Display")
add_textbox(slide, Inches(9.6), Inches(3.75), Inches(2.55), Inches(1.4), "What this receiver
makes possible for
this sender, in this task", 18, TEXT_DARK)

for x in [Inches(4.0), Inches(8.3)]:
    arrow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, x, Inches(3.72), Inches(0.55), Inches(0.5))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = ACCENT
    arrow.line.color.rgb = ACCENT

add_textbox(slide, Inches(0.78), Inches(6.2), Inches(11.8), Inches(0.55), "Reasoning can remain continuous and private; the interface becomes discrete because interaction requires public references.", 18, TEXT_DARK)


# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = BG_LIGHT

add_textbox(slide, Inches(0.7), Inches(0.55), Inches(8.8), Inches(0.65), "Interaction Shapes the Interface", 28, TEXT_DARK, bold=True, font="Aptos Display")
add_textbox(slide, Inches(0.72), Inches(1.2), Inches(9.4), Inches(0.72), "Agents must act under partial access to each other's latent capabilities. What they can use is only what becomes visible through the interface.", 17, MUTED)

add_round_box(slide, Inches(0.75), Inches(2.0), Inches(3.8), Inches(3.95), CARD, RGBColor(0xE5, 0xD8, 0xCB))
add_textbox(slide, Inches(1.0), Inches(2.25), Inches(3.1), Inches(0.4), "Constraint", 22, TEXT_DARK, bold=True, font="Aptos Display")
add_bullet_list(slide, Inches(1.0), Inches(2.8), Inches(3.15), Inches(2.6), [
    "No agent sees the full latent space of the other.",
    "Ability must be inferred from exposed signals and handles.",
    "Affordances are perceived through the interface, not directly.",
], size=18)

add_round_box(slide, Inches(4.8), Inches(2.0), Inches(3.8), Inches(3.95), RGBColor(0xF0, 0xDE, 0xD6), RGBColor(0xE5, 0xC0, 0xAF))
add_textbox(slide, Inches(5.05), Inches(2.25), Inches(3.1), Inches(0.4), "Interaction", 22, TEXT_DARK, bold=True, font="Aptos Display")
add_bullet_list(slide, Inches(5.05), Inches(2.8), Inches(3.15), Inches(2.6), [
    "Agents adapt what they expose to what they expect the partner will need.",
    "Selection and execution are one example of this pressure.",
    "Communication shifts toward usable and interpretable references.",
], size=18)

add_round_box(slide, Inches(8.85), Inches(2.0), Inches(3.75), Inches(3.95), RGBColor(0xD8, 0xE6, 0xD8), RGBColor(0xB6, 0xCF, 0xB6))
add_textbox(slide, Inches(9.1), Inches(2.25), Inches(3.0), Inches(0.4), "Consequence", 22, TEXT_DARK, bold=True, font="Aptos Display")
add_bullet_list(slide, Inches(9.1), Inches(2.8), Inches(3.05), Inches(2.6), [
    "Interface and latent space remain distinct.",
    "The interface is shaped by task-specific affordances.",
    "Language may evolve toward compact, partner-usable forms.",
], size=18)

footer = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(6.72), Inches(13.333), Inches(0.78))
footer.fill.solid()
footer.fill.fore_color.rgb = BG_DARK
footer.line.color.rgb = BG_DARK
add_textbox(slide, Inches(0.75), Inches(6.93), Inches(12), Inches(0.3), "Core claim: interaction shapes the interface because agents must expose what matters without revealing the full latent space.", 16, TEXT_LIGHT)

prs.save("docs/Language_as_Interface_Concept_Slides.pptx")
